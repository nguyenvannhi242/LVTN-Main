from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "report_updates"

BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x59, 0x59, 0x59)

def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    # Set to A4 size (standard for Vietnamese thesis)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)      # 2.54 cm
    section.bottom_margin = Inches(1.0)   # 2.54 cm
    section.left_margin = Inches(1.25)    # 3.17 cm (binding margin)
    section.right_margin = Inches(0.79)   # 2.0 cm

    # Configure Normal Style (Times New Roman, 12pt, 1.25 line spacing, Justify)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Configure Heading Styles (Times New Roman, Bold, Black, space before/after)
    for name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Add custom tight style if needed
    if "Body Tight" not in doc.styles:
        style = doc.styles.add_style("Body Tight", 1) # 1 is WD_STYLE_TYPE.PARAGRAPH
        style.base_style = normal
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("BÁO CÁO LUẬN VĂN TỐT NGHIỆP")
    run.bold = True
    run.font.size = Pt(14)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("CHƯƠNG 1: MỞ ĐẦU")
    run.bold = True
    run.font.size = Pt(20)
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    run = p.add_run("Bản cập nhật biên soạn mới nhất")
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = GRAY

    metadata = (
        ("Đề tài chính thức", "Nghiên cứu giải pháp phát hiện lạm dụng Trí tuệ nhân tạo trong bài tập lập trình C++ bằng mô hình học sâu và học máy giải thích được"),
        ("Đối chiếu kỹ thuật", "GraphCodeBERT, LightGBM + SHAP, LangGraph Agent, Django + FastAPI, Medallion Architecture"),
        ("Quy chuẩn định dạng", "Times New Roman 12pt, Giãn dòng 1.25, Căn lề đóng gáy (Justified)"),
        ("Ngày cập nhật", "29/05/2026"),
    )
    
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r1 = p.add_run(f"{label}: ")
        r1.bold = True
        p.add_run(value)
        
    doc.add_section(WD_SECTION_START.NEW_PAGE)

def build_chapter_1() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_doc(doc)
    add_title_page(doc)

    # Content structure
    content = [
        ("CHƯƠNG 1: MỞ ĐẦU", 1),
        ("Chương này dẫn dắt vào đề tài nghiên cứu, trình bày lý do lựa chọn đề tài thông qua việc phân tích bối cảnh công nghệ và tính liêm chính học thuật trong giáo dục CNTT. Đồng thời, chương cũng đặc tả mục tiêu nghiên cứu, đối tượng và phạm vi nghiên cứu, các phương pháp tiếp cận khoa học được áp dụng, ý nghĩa đóng góp thực tiễn và cấu trúc tổng thể của luận văn.", 0),
        
        ("1.1. Lý do chọn đề tài", 2),
        ("Trong kỷ nguyên số, sự phát triển bùng nổ của Trí tuệ nhân tạo (AI) và các Mô hình ngôn ngữ lớn (LLM) như ChatGPT (OpenAI), Claude (Anthropic), Gemini (Google) và Qwen (Alibaba) đã định hình lại nhiều khía cạnh của đời sống kinh tế - xã hội. Trong lĩnh vực kỹ nghệ phần mềm, AI đã chứng minh năng lực vượt trội trong việc tự động sinh mã nguồn, tối ưu hóa thuật toán, giải thích logic mã và phát hiện lỗi cú pháp. Các công cụ hỗ trợ lập trình như GitHub Copilot hay Cursor đã trở thành trợ thủ đắc lực giúp lập trình viên chuyên nghiệp nâng cao năng suất lao động từ 20% đến 50%.", 0),
        ("Tuy nhiên, sự phổ biến rộng rãi và dễ dàng tiếp cận của các công cụ sinh mã tự động này cũng tạo ra những hệ lụy tiêu cực chưa từng có đối với môi trường giáo dục đại học, đặc biệt là các học phần lập trình cơ sở và nâng cao thuộc khối ngành Công nghệ thông tin (CNTT). Sinh viên hoàn toàn có thể sử dụng AI để hoàn thành các bài tập lập trình, đồ án môn học hoặc thậm chí là các bài thi thực hành trực tuyến chỉ bằng cách sao chép đề bài và dán vào ô trò chuyện của LLM. Hành vi lạm dụng này nếu không được kiểm soát sẽ làm suy giảm nghiêm trọng năng lực tư duy thuật toán, kỹ năng tự giải quyết vấn đề và chất lượng chuyên môn thực tế của sinh viên khi ra trường, gây ảnh hưởng trực tiếp đến chất lượng nguồn nhân lực phần mềm quốc gia.", 0),
        ("Trước khi làn sóng Generative AI xuất hiện, bài toán giám sát học thuật đối với các bài tập lập trình chủ yếu dựa trên các công cụ phát hiện đạo văn truyền thống như MOSS (Measure Of Software Similarity) của Đại học Stanford hoặc JPlag. Nguyên lý hoạt động của các công cụ này là phân tích tính tương đồng về mặt cấu trúc văn bản hoặc biểu diễn cây cú pháp trừu tượng (Abstract Syntax Tree - AST) giữa các bài nộp của sinh viên trong cùng một lớp. Tuy nhiên, các công cụ này hoàn toàn \"bất lực\" trước mã nguồn do AI sinh ra. Bản chất của các mô hình LLM là sinh từ ngẫu nhiên dựa trên phân phối xác suất ngôn ngữ lớn, cho phép chúng tạo ra hàng chục phiên bản mã nguồn khác nhau về mặt cú pháp biểu diễn nhưng giống hệt nhau về mặt logic nghiệp vụ cho cùng một yêu cầu bài toán. Đoạn mã do AI sinh ra là duy nhất, không trùng khớp với bất kỳ bài nộp nào khác hoặc bất kỳ đoạn mã có sẵn nào trên mạng Internet, giúp sinh viên dễ dàng vượt qua các bộ lọc đạo văn truyền thống.", 0),
        ("Từ những thách thức trên, nhu cầu nghiên cứu một giải pháp kỹ thuật có khả năng tự động nhận diện xem một đoạn mã nguồn lập trình cụ thể là do con người tự viết hay do AI tạo ra (AI-generated code detection) trở nên cấp thiết hơn bao giờ hết. Hơn thế nữa, trong bối cảnh sư phạm, hệ thống kiểm soát không thể chỉ đưa ra một kết luận nhị phân dạng \"hộp đen\" (black-box) với một điểm số tin cậy chung chung. Giảng viên cần những bằng chứng rõ ràng, mang tính trực quan và giải thích được (Explainable AI - XAI) về các đặc trưng phong cách lập trình (Stylometry) hoặc dấu vết cấu trúc để thuyết phục sinh viên và đưa ra quyết định xử lý khách quan. Đồng thời, giải pháp này cần được tích hợp vào một hệ thống phần mềm có kiến trúc vững chắc, có khả năng xử lý bất đồng bộ thời gian thực và giao diện thân thiện để sẵn sàng triển khai thực tế tại các cơ sở đào tạo.", 0),
        ("Nhận thức được tầm quan trọng và tính cấp thiết đó, nhóm nghiên cứu đã lựa chọn thực hiện đề tài: \"Nghiên cứu giải pháp phát hiện lạm dụng Trí tuệ nhân tạo trong bài tập lập trình C++ bằng mô hình học sâu và học máy giải thích được\". Đề tài hướng tới việc xây dựng một hệ thống phát hiện toàn diện, kết hợp hài hòa giữa năng lực biểu diễn ngữ nghĩa sâu sắc của mạng học sâu và tính minh bạch, giải thích được của mô hình học máy dạng bảng dựa trên đặc trưng phong cách.", 0),
        
        ("1.2. Mục tiêu nghiên cứu", 2),
        ("1.2.1. Mục tiêu tổng quát", 3),
        ("Mục tiêu tổng quát của đề tài là nghiên cứu, thiết kế và xây dựng thành công hệ thống C++ AI Code Detector có độ chính xác cao, có khả năng giải thích trực quan và kết cấu kiến trúc phần mềm tối ưu để ứng dụng phát hiện hành vi lạm dụng AI trong bài tập lập trình C++ của sinh viên tại các trường đại học.", 0),
        ("1.2.2. Mục tiêu cụ thể", 3),
        ("Để đạt được mục tiêu tổng quát, đề tài đề ra 4 mục tiêu cụ thể sau:", 0),
        ("1. Xây dựng tập dữ liệu C++ đối ngẫu quy mô lớn và chuẩn hóa: Thiết lập quy trình thu thập dữ liệu tự động từ GitHub với các điều kiện lọc thời gian nghiêm ngặt (trước năm 2022) để thu thập mã nguồn con người viết sạch, kết hợp sinh phiên bản đối ngẫu AI tương ứng bằng LLM API; thực hiện làm sạch dữ liệu và lọc trùng lặp tuyệt đối.", 0),
        ("2. Huấn luyện và tối ưu hóa Mô hình lai (Hybrid Model):\n  • Huấn luyện mô hình học sâu GraphCodeBERT để khai thác ngữ nghĩa cấu trúc thông qua đồ thị luồng dữ liệu (Data Flow Graph) của chương trình C++.\n  • Huấn luyện mô hình học máy LightGBM trên bộ 32 đặc trưng tĩnh trích xuất từ phong cách lập trình (Stylometry), bố cục, quy ước đặt tên và các số liệu đo lường phần mềm.\n  • Thiết lập cơ chế định tuyến (Router Node) phân tách phong cách viết Basic và OOP để tối ưu hóa việc gọi mô hình chuyên biệt.", 0),
        ("3. Tích hợp cơ chế giải thích hậu nghiệm (XAI) và lập luận tác nhân:\n  • Áp dụng thuật toán TreeSHAP để giải thích định lượng đóng góp của từng đặc trưng phong cách tĩnh vào quyết định phân loại của LightGBM.\n  • Xây dựng hệ thống tác nhân thông minh LangGraph Agent để điều phối quy trình phân tích và sử dụng LLM sinh báo cáo định tính bằng Tiếng Việt tự nhiên dễ hiểu cho giảng viên.", 0),
        ("4. Phát triển hệ thống phần mềm doanh nghiệp hoàn chỉnh: Hiện thực hóa giải pháp dưới dạng ứng dụng Web phân rã microservices sử dụng Django làm Web Portal quản lý dữ liệu đa tầng (Medallion Architecture) và FastAPI làm AI Serving Orchestrator hỗ trợ truyền tải tiến trình thời gian thực qua Server-Sent Events (SSE).", 0),
        
        ("1.3. Đối tượng và phạm vi nghiên cứu", 2),
        ("1.3.1. Đối tượng nghiên cứu", 3),
        ("• Đối tượng nghiên cứu chính: Mã nguồn (source code) của các chương trình máy tính được viết bằng ngôn ngữ lập trình C++ (bao gồm các chuẩn C++ từ cổ điển như C++98 đến hiện đại C++20).\n• Đối tượng nghiên cứu công nghệ: Các mô hình ngôn ngữ lớn chuyên lập trình, kiến trúc mạng Transformer, cơ chế Attention, các kỹ thuật học sâu biểu diễn mã nguồn (GraphCodeBERT), các thuật toán học máy dạng bảng (LightGBM, XGBoost, Random Forest), lý thuyết trò chơi hợp tác (Shapley values/SHAP) và mô hình tác nhân thông minh (AI Agent).", 0),
        ("1.3.2. Phạm vi nghiên cứu", 3),
        ("• Phạm vi bài tập lập trình: Giới hạn trong khuôn khổ các bài tập thực hành lập trình của sinh viên CNTT, được chia làm hai phong cách viết chính:\n  - Basic C++: Các bài toán thuật toán tuyến tính, cấu trúc điều khiển cơ bản, mảng, hàm đơn giản.\n  - OOP C++: Các chương trình hướng đối tượng có cấu trúc phức tạp hơn, sử dụng lớp (class), kế thừa (inheritance), đa hình (polymorphism), khuôn mẫu (template) và cấu trúc đa tệp.\n• Phạm vi công nghệ giải thích: Tập trung vào giải thích cấp độ đặc trưng phong cách lập trình tĩnh (Feature-level Explainability) bằng phương pháp SHAP trên mô hình LightGBM. Các phương pháp giải thích cấp độ token (như Layer Integrated Gradients) nằm ngoài phạm vi đóng góp chính của đề tài do hạn chế về tính ứng dụng thực tiễn sư phạm.\n• Phạm vi hệ thống phần mềm: Triển khai theo mô hình Hybrid Local-Cloud (Django chạy local quản lý nghiệp vụ và cơ sở dữ liệu PostgreSQL; FastAPI điều phối pipeline AI; GPU runtime xử lý inference nặng chạy local hoặc cloud Colab).", 0),
        
        ("1.4. Phương pháp nghiên cứu", 2),
        ("Để thực hiện đề tài này, nhóm nghiên cứu kết hợp chặt chẽ giữa nghiên cứu lý thuyết và nghiên cứu thực nghiệm thực tiễn:", 0),
        ("1.4.1. Phương pháp nghiên cứu lý thuyết", 3),
        ("• Nghiên cứu tài liệu: Khảo sát các công trình khoa học uy tín trong nước và quốc tế liên quan đến bài toán phát hiện tác giả mã nguồn (Code Authorship Attribution) và phát hiện mã nguồn do AI sinh ra.\n• Phân tích thuật toán sâu: Tìm hiểu cơ sở toán học của kiến trúc Transformer và mô hình GraphCodeBERT; nghiên cứu lý thuyết trò chơi hợp tác của Lloyd Shapley và cách áp dụng tính toán đóng góp biên thông qua TreeSHAP; nghiên cứu cơ chế tính toán độ hoang mang ngôn ngữ (Perplexity) trên phân phối xác suất token.", 0),
        ("1.4.2. Phương pháp nghiên cứu thực nghiệm", 3),
        ("• Xây dựng Dataset: Thu thập dữ liệu thực tế từ GitHub, viết công cụ lọc trùng lặp và làm sạch code; viết prompt gọi API của GPT-4o-mini để tạo dữ liệu đối ngẫu tương ứng; gán nhãn tự động xây dựng tập dữ liệu 41,999 mẫu.\n• Phát triển công cụ trích xuất đặc trưng: Thiết kế và hiện thực hóa lớp CppFeatureExtractorV8 trích xuất chính xác 32 thuộc tính phong cách tĩnh từ tệp tin C++.\n• Huấn luyện và Tối ưu hóa mô hình: Sử dụng PyTorch và Hugging Face để fine-tune GraphCodeBERT với kỹ thuật Sliding Window và Stratified Group K-Fold; sử dụng thư viện LightGBM để huấn luyện bộ phân loại tĩnh; thực hiện tối ưu hóa siêu tham số.\n• Kiểm thử và đánh giá: Thiết lập các kịch bản kiểm thử In-Distribution (unified_test) và Out-Of-Distribution (OOD trên 10 mô hình ngôn ngữ khác nhau bao gồm DeepSeek-R1 và o3-mini); sử dụng các độ đo Accuracy, Precision, Recall, và F1-Macro để đánh giá hiệu năng.\n• Xây dựng phần mềm: Lập trình ứng dụng Django và FastAPI kết nối qua hàng Celery Workers và Redis Broker; kiểm tra hoạt động truyền dẫn SSE và lưu trữ Medallion.", 0),
        
        ("1.5. Ý nghĩa khoa học và thực tiễn của đề tài", 2),
        ("1.5.1. Ý nghĩa khoa học", 3),
        ("• Đề xuất mô hình định tuyến lai đột phá: Đóng góp một giải pháp phân loại lai ghép (Hybrid) kết hợp tối ưu giữa \"ngữ nghĩa cấu trúc sâu\" (deep representation) và \"vân tay phong cách tĩnh\" (stylometry), khắc phục triệt để hạn chế của từng phương pháp khi chạy độc lập.\n• Chứng minh tính ưu việt của Feature-level XAI: Đóng góp một nghiên cứu thực nghiệm chứng minh rằng việc giải thích ở cấp độ thuộc tính phong cách phần mềm có giá trị diễn giải đối với con người cao hơn nhiều so với việc trực quan hóa heatmap cấp độ token thô rời rạc.\n• Thiết lập quy trình đánh giá độ bền vững ngoại miền phân phối: Xây dựng một benchmark kiểm thử OOD toàn diện trên 10 mô hình ngôn ngữ lớn thế hệ mới, chứng minh khả năng tổng quát hóa bền vững của mô hình lai trước các LLM suy luận logic sâu.", 0),
        ("1.5.2. Ý nghĩa thực tiễn", 3),
        ("• Bảo vệ liêm chính học thuật: Cung cấp cho các trường đại học một công cụ đắc lực, tự động để phát hiện và ngăn chặn tình trạng gian lận, lạm dụng các công cụ sinh mã tự động của sinh viên trong quá trình học tập và kiểm tra lập trình C++.\n• Hỗ trợ giảng dạy và đánh giá: Báo cáo định tính chi tiết từ Critique Node giúp giảng viên hiểu rõ thói quen viết code của sinh viên, từ đó đưa ra các phản hồi sư phạm thích hợp hoặc phát hiện sinh viên cần hỗ trợ học tập.\n• Kiến trúc phần mềm sẵn sàng triển khai: Sản phẩm phần mềm với cấu trúc Django + FastAPI phân rã microservice, sử dụng lưu trữ Medallion và tối ưu hóa nạp/giải phóng bộ nhớ đồ họa có khả năng vận hành ổn định trong môi trường máy chủ nhà trường với chi phí phần cứng tối thiểu.", 0),
        
        ("1.6. Cấu trúc của luận văn", 2),
        ("Luận văn được tổ chức thành 5 chương chính với nội dung phân bổ logic như sau:\n• CHƯƠNG 1: MỞ ĐẦU: Trình bày lý do chọn đề tài, mục tiêu, đối tượng, phạm vi nghiên cứu, phương pháp tiếp cận khoa học, ý nghĩa khoa học - thực tiễn và cấu trúc của luận văn.\n• CHƯƠNG 2: NỀN TẢNG LÝ THUYẾT VÀ CÔNG TRÌNH LIÊN QUAN: Trình bày cơ sở lý thuyết về kiến trúc Transformer, GraphCodeBERT, lý thuyết trò chơi SHAP, Perplexity; phân tích, so sánh các nghiên cứu liên quan trong nước và quốc tế để chỉ ra các khoảng trống nghiên cứu (Research Gaps).\n• CHƯƠNG 3: GIẢI PHÁP ĐỀ XUẤT: Đặc tả chi tiết giải pháp đề xuất bao gồm: sơ đồ kiến trúc tổng thể, quy trình xây dựng tập dữ liệu đối ngẫu C++, bộ trích xuất tĩnh CppFeatureExtractorV8 với 32 đặc trưng, phương pháp huấn luyện mô hình GraphCodeBERT/LightGBM, thiết kế đồ thị tác nhân LangGraph và đặc tả giao diện hệ thống Django/FastAPI.\n• CHƯƠNG 4: THỰC NGHIỆM VÀ ĐÁNH GIÁ KẾT QUẢ: Mô tả môi trường thực nghiệm, phân bố tập dữ liệu; trình bày chi tiết kết quả của 5 kịch bản thực nghiệm đối chứng (ML, XAI bằng SHAP, DL chuyên biệt, mô hình Hybrid đề xuất, và đánh giá OOD Generalization); đưa ra các phân tích và nhận định học thuật.\n• CHƯƠNG 5: KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN: Tổng kết lại các kết quả khoa học và thực tế đã đạt được, thẳng thắn chỉ ra các hạn chế kỹ thuật hiện tại của hệ thống và đề xuất các hướng nghiên cứu tiềm năng trong tương lai.", 0),
    ]

    for text, level in content:
        if level == 0:
            doc.add_paragraph(text)
        else:
            doc.add_paragraph(text, style=f"Heading {level}")

    out_path = OUT_DIR / "Chuong_1_cap_nhat.docx"
    doc.save(out_path)
    print(f"Successfully generated docx at: {out_path}")

if __name__ == "__main__":
    build_chapter_1()
