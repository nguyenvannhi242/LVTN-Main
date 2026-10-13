from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "report_updates"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
GRAY = RGBColor(0x66, 0x66, 0x66)
BLACK = RGBColor(0x00, 0x00, 0x00)


@dataclass
class TableSpec:
    insert_after: str
    title: str
    headers: list[str]
    rows: list[list[str]]
    widths: list[float]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15

    if "Body Tight" not in doc.styles:
        style = doc.styles.add_style("Body Tight", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def add_title_page(doc: Document, chapter_title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BÁO CÁO LUẬN VĂN TỐT NGHIỆP")
    run.bold = True
    run.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(chapter_title)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = BLACK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(subtitle)
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = GRAY

    for label, value in (
        ("Đề tài", "Xây dựng hệ thống phát hiện mã nguồn C++ do AI sinh ra"),
        ("Phiên bản", "Bản cập nhật v2 - giữ chuẩn học thuật, chỉnh đúng phạm vi hiện tại"),
        ("Nguồn đối chiếu", "Chuong 2+3.docx, README.md, ARCHITECTURE_MAP.md, src/"),
        ("Ngày tạo", "20/05/2026"),
    ):
        p = doc.add_paragraph(style="Body Tight")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p.add_run(f"{label}: ")
        r1.bold = True
        p.add_run(value)

    doc.add_section(WD_SECTION_START.NEW_PAGE)


def add_paragraph(doc: Document, text: str, style: str = "Normal") -> None:
    doc.add_paragraph(text, style=style)


def add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def add_table(doc: Document, spec: TableSpec) -> None:
    title = doc.add_paragraph(style="Body Tight")
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(spec.title)
    run.bold = True

    table = doc.add_table(rows=1, cols=len(spec.headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    for i, header in enumerate(spec.headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(spec.widths[i])
        cell.text = header
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            run.bold = True
            run.font.size = Pt(10)

    for row_values in spec.rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.width = Inches(spec.widths[i])
            cell.text = value
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in para.runs:
                run.font.size = Pt(10)
    doc.add_paragraph("", style="Body Tight")


def ch2() -> tuple[str, str, list[tuple[str, int]], list[TableSpec]]:
    parts = [
        ("CHƯƠNG 2: NỀN TẢNG LÝ THUYẾT VÀ CÔNG TRÌNH LIÊN QUAN", 1),
        ("Chương này trình bày các nền tảng lý thuyết cốt lõi làm cơ sở cho đề tài \"Xây dựng hệ thống phát hiện AI-generated code với GraphCodeBERT, Agent và Explainable AI\". Nội dung bao gồm: bài toán phát hiện mã nguồn do AI tạo ra trong bối cảnh học thuật, các mô hình ngôn ngữ tiền huấn luyện cho mã nguồn từ BERT, RoBERTa, CodeBERT đến GraphCodeBERT, các kỹ thuật giải thích mô hình đang được sử dụng trong phạm vi hiện tại của đề tài, và tổng hợp các công trình nghiên cứu liên quan.", 0),
        ("2.1 Cơ sở lý thuyết", 2),
        ("2.1.1 Bài toán phát hiện mã nguồn do AI tạo ra", 3),
        ("Sự bùng nổ của các mô hình ngôn ngữ lớn như ChatGPT (OpenAI, 2022), GitHub Copilot (2021), và Gemini (Google, 2024) đã tạo ra một bước ngoặt trong lĩnh vực công nghệ phần mềm. Các hệ thống này có khả năng sinh mã nguồn hoàn chỉnh từ mô tả bằng ngôn ngữ tự nhiên, tái cấu trúc mã, sửa lỗi và tạo tài liệu kỹ thuật với chất lượng ngày càng gần với sản phẩm do con người viết.", 0),
        ("Trong bối cảnh giáo dục đại học, đặc biệt là các học phần lập trình, việc sinh viên sử dụng AI để hoàn thành bài tập hoặc đồ án mà không tự viết code đặt ra thách thức trực tiếp với tính liêm chính học thuật [27]. Các công cụ phát hiện đạo văn truyền thống như MOSS chủ yếu dựa trên so khớp chuỗi hoặc cấu trúc tương đồng giữa các bài nộp, nên khó phát hiện mã do AI sinh ra khi mã đó hợp lệ về cú pháp nhưng không sao chép trực tiếp từ một nguồn cụ thể.", 0),
        ("Bài toán AI Code Detection được định nghĩa như sau: cho một đoạn mã nguồn C++ bất kỳ, hệ thống cần phân loại đoạn mã đó thuộc một trong hai nhóm: do con người viết hoặc do AI tạo ra. Đây là bài toán phân loại nhị phân trên dữ liệu dạng mã nguồn. Điểm khó nằm ở chỗ nhãn AI/Human không chỉ phụ thuộc vào cú pháp đúng sai, mà phụ thuộc vào dấu vết phong cách, cấu trúc, xác suất ngôn ngữ và thói quen trình bày mã.", 0),
        ("2.1.2 Kiến trúc Transformer và các mô hình ngôn ngữ cho mã nguồn", 3),
        ("Transformer (Vaswani et al., 2017) [21] là kiến trúc mạng nơ-ron nền tảng cho hầu hết các mô hình ngôn ngữ hiện đại. Điểm đột phá của Transformer so với RNN và LSTM nằm ở cơ chế Self-Attention, cho phép mô hình xem xét đồng thời quan hệ giữa tất cả vị trí trong chuỗi đầu vào mà không bị giới hạn bởi khoảng cách tuần tự.", 0),
        ("Cơ chế Multi-Head Self-Attention tính ba loại vector cho mỗi token đầu vào: Query (Q), Key (K), và Value (V). Điểm Attention được tính theo công thức:", 0),
        ("Attention(Q, K, V) = softmax(QK^T / √d_k) × V", 0),
        ("trong đó d_k là số chiều của vector Key, đóng vai trò hệ số chuẩn hóa để ổn định gradient. Multi-Head Attention chạy song song nhiều phép Attention trên các không gian con khác nhau, sau đó ghép kết quả lại, giúp mô hình nắm bắt nhiều loại quan hệ ngữ nghĩa đồng thời [21].", 0),
        ("BERT (Devlin et al., 2019) là mô hình tiền huấn luyện sử dụng phần Encoder của Transformer. BERT được huấn luyện trên Masked Language Modeling và Next Sentence Prediction. RoBERTa (Liu et al., 2019) cải tiến BERT bằng cách loại bỏ NSP, huấn luyện lâu hơn và trên tập dữ liệu lớn hơn, trở thành nền tảng cho nhiều mô hình chuyên biệt về mã nguồn.", 0),
        ("CodeBERT (Feng et al., 2020) [22] là mô hình tiền huấn luyện lưỡng thức cho cặp ngôn ngữ lập trình và ngôn ngữ tự nhiên. Mô hình kế thừa kiến trúc RoBERTa và được huấn luyện trên CodeSearchNet với hai nhiệm vụ chính: Masked Language Modeling và Replaced Token Detection. Hạn chế của CodeBERT là chủ yếu xử lý mã nguồn ở dạng chuỗi token, chưa khai thác đầy đủ cấu trúc ngữ nghĩa của chương trình.", 0),
        ("GraphCodeBERT (Guo et al., 2020) [23] khắc phục hạn chế này bằng cách tích hợp Data Flow Graph vào quá trình tiền huấn luyện. DFG biểu diễn quan hệ giữa biến và các vị trí sử dụng biến, giúp mô hình học được mối liên hệ ngữ nghĩa của chương trình thay vì chỉ dựa vào thứ tự token. GraphCodeBERT bổ sung graph-guided masked attention và edge prediction, qua đó đạt kết quả tốt hơn CodeBERT trên các tác vụ như code search, clone detection, code translation và code refinement [23].", 0),
        ("Trong hệ thống của nhóm, GraphCodeBERT (microsoft/graphcodebert-base) được chọn làm mô hình phân loại lõi vì khả năng nhận thức cấu trúc ngữ nghĩa của mã C++, đặc biệt khi phân biệt mã Basic và mã OOP. Mô hình được fine-tune với lớp phân loại nhị phân trên tập dữ liệu C++ tự xây dựng.", 0),
        ("2.1.3 Phân tích có tính giải thích", 3),
        ("Trong bối cảnh mô hình học sâu thường hoạt động như hộp đen, Explainable AI (XAI) là yêu cầu cần thiết để người dùng hiểu vì sao hệ thống kết luận một đoạn code là AI-generated hoặc Human-written. Theo khảo sát [17], các phương pháp XAI có thể chia thành hai nhóm: interpretability nội tại và interpretability hậu nghiệm.", 0),
        ("Ở phạm vi hiện tại của đề tài, kỹ thuật giải thích chính được sử dụng là SHAP (Shapley Additive Explanations) cho mô hình đặc trưng LightGBM. SHAP dựa trên Shapley values trong lý thuyết trò chơi, phân bổ đóng góp của từng đặc trưng vào đầu ra dự đoán. Với một mô hình f và tập đặc trưng F, giá trị Shapley của đặc trưng i có thể hiểu là trung bình đóng góp biên của đặc trưng đó trên mọi tập con đặc trưng có thể có [18].", 0),
        ("Việc sử dụng SHAP phù hợp với hướng triển khai hiện tại vì hệ thống có nhánh phân tích đặc trưng tĩnh của mã nguồn như độ dài dòng, tỉ lệ comment, độ nhất quán dấu ngoặc, độ dài định danh, entropy khoảng trắng, độ phức tạp cyclomatic và các chỉ số Halstead. Khi LightGBM đưa ra xác suất AI, SHAP giúp chỉ ra đặc trưng nào kéo quyết định về phía AI hoặc Human, từ đó tạo ra lời giải thích có thể đọc được bởi giảng viên.", 0),
        ("Vì phiên bản báo cáo hiện tại ưu tiên hướng giải thích feature-level bằng SHAP/LightGBM, các kỹ thuật giải thích token-level thử nghiệm không được trình bày như đóng góp chính. Cách viết này giúp phần lý thuyết bám sát đúng thành phần đang dùng để diễn giải kết quả trong hệ thống.", 0),
        ("2.1.4 Perplexity và tín hiệu xác suất ngôn ngữ", 3),
        ("Perplexity (PPL) là thước đo mức độ ngạc nhiên của một mô hình ngôn ngữ khi đọc một chuỗi token. Nếu L là cross-entropy loss trung bình trên toàn bộ chuỗi đầu vào, Perplexity được tính theo công thức:", 0),
        ("PPL = exp(L)", 0),
        ("PPL thấp nghĩa là mô hình ít ngạc nhiên, tức chuỗi mã dễ dự đoán hơn; PPL cao nghĩa là chuỗi mã khó dự đoán hơn. Giả thuyết nền tảng của đề tài là mã do AI sinh ra thường tuân theo các pattern phổ biến và được chuẩn hóa tốt, nên có thể có PPL thấp hơn mã do con người viết. Tuy nhiên, PPL không nên dùng độc lập vì mã ngắn, mã theo template hoặc mã bài tập cơ bản của sinh viên cũng có thể dễ dự đoán.", 0),
        ("Trong hệ thống, PPL đóng vai trò tín hiệu bổ trợ cho bước đánh giá kết quả. Khi AI Score và PPL mâu thuẫn, ví dụ AI Score cao nhưng PPL cũng cao, hệ thống có thể xem kết quả là nhập nhằng và cần đánh giá lại bằng nhánh mô hình khác.", 0),
        ("2.1.5 Chain-of-Thought Prompting và nhận xét bằng LLM", 3),
        ("Chain-of-Thought Prompting (Wei et al., 2022) [16] là kỹ thuật khuyến khích mô hình ngôn ngữ sinh ra các bước lập luận trung gian trước khi đưa ra kết luận. Trong đề tài, nguyên lý này được áp dụng ở mức sinh nhận xét: LLM nhận các tín hiệu như AI Score, đặc trưng SHAP nổi bật, perplexity và đoạn code mẫu, sau đó tạo phần giải thích kỹ thuật thay vì chỉ trả về nhãn AI/Human.", 0),
        ("2.2 Các công trình nghiên cứu liên quan", 2),
        ("2.2.1 Các nghiên cứu về phát hiện mã do AI sinh ra", 3),
        ("Krish và Sotiriades (2025) [1] nghiên cứu trực tiếp bài toán xác định mã do ChatGPT tạo ra so với mã do con người viết. Nhóm tác giả sử dụng đặc trưng từ mô hình pre-trained kết hợp CNN để phân loại, đồng thời đề xuất SCAE để xử lý trường hợp code bị biến đổi. Hạn chế chính là nghiên cứu tập trung vào ChatGPT và chưa tích hợp cơ chế giải thích kết quả.", 0),
        ("Một nghiên cứu năm 2024 [15] đạt khoảng 98% accuracy trong phân biệt code ChatGPT với code người viết bằng cách kết hợp embedding features với DNN, Random Forest, XGBoost và white-box features với Bayes classifier. Công trình này có đề cập đến cả phân loại và giải thích, nhưng giải thích vẫn chủ yếu dừng ở mức feature importance, chưa tạo thành một hệ thống phân tích đầy đủ cho người dùng cuối.", 0),
        ("2.2.2 Các nghiên cứu về nhận diện tác giả mã nguồn", 3),
        ("Nhận diện tác giả mã nguồn là hướng nghiên cứu có liên hệ chặt với AI code detection vì cả hai đều khai thác dấu vết phong cách trong mã. Kothari et al. (2007) [2] dùng Bayes classifier với 4-character n-gram, đạt 61% accuracy cho 12 tác giả. Yang et al. (2020) [10] kết hợp character n-gram, TF-IDF và Logistic Regression, đạt 94.28% trên tập AI-SOCO.", 0),
        ("Ở hướng học sâu, Abuhamad et al. (2019) [3] dùng CNN kết hợp TF-IDF và word embedding, đạt 99.4% accuracy trên Google Code Jam. Abuhamad et al. (2021) [4] tiếp tục phát triển với LSTM/RNN kết hợp Random Forest, đạt 96.34% trên C++ và Python. Multi-χ (Abuhamad et al., 2020) [9] xử lý bài toán đa tác giả trong một file code bằng cách chia code thành các segment và dùng RNN để xác nhận hoặc phủ nhận cùng tác giả.", 0),
        ("Các phương pháp Transformer cũng được áp dụng cho authorship attribution. Mahbub et al. (2022) [12] dùng Transformer ở cấp segment-level. AuthAttLyzer-V2 (2024) [31] phân tích trên 24.000 file C++ từ 3.000 tác giả. CLAVE (Álvarez-Fidalgo et al., 2025) [5][7] dùng contrastive learning với Transformer encoder cho authorship verification, đạt 92.3% accuracy chỉ với 6 file mỗi lập trình viên.", 0),
        ("2.2.3 Các mô hình biểu diễn mã nguồn và benchmark", 3),
        ("Ngoài CodeBERT [22] và GraphCodeBERT [23], một số kiến trúc biểu diễn mã nguồn đáng chú ý gồm SeTransformer [24], kết hợp code tokens và AST, và Deep Graph CNN [13], sử dụng graph convolution trên AST và Control Flow Graph. Về benchmark, CodeXGLUE (Lu et al., 2021) [30] cung cấp 14 tập dữ liệu cho 10 tác vụ hiểu mã nguồn, nhưng chưa có tập chuyên biệt cho bài toán phân loại AI vs Human trên C++.", 0),
        ("2.2.4 Khoảng trống nghiên cứu", 3),
        ("Từ Bảng 2.3 có thể rút ra bốn khoảng trống chính. Thứ nhất, còn thiếu tập dữ liệu C++ chuyên biệt cho AI Code Detection, đặc biệt là dữ liệu phân biệt Basic/OOP và có kiểm soát thời điểm trước khi LLM phổ biến. Thứ hai, nhiều công trình chỉ đưa ra nhãn phân loại hoặc confidence score mà chưa cung cấp giải thích đủ rõ cho người dùng. Thứ ba, các nghiên cứu thường tập trung vào một tín hiệu đơn lẻ, trong khi bài toán thực tế cần kết hợp mô hình học sâu, đặc trưng tĩnh, perplexity và nhận xét LLM. Thứ tư, các hệ thống nghiên cứu thường chưa gắn chặt với một ứng dụng có giao diện, lưu trữ kết quả và khả năng truy xuất lại phân tích.", 0),
        ("2.3 Tổng kết chương", 2),
        ("Chương 2 đã trình bày nền tảng lý thuyết và các công trình liên quan cho đề tài. Điểm cần nhấn mạnh là giải pháp không chỉ dựa vào một mô hình phân loại, mà kết hợp GraphCodeBERT cho biểu diễn mã nguồn, LightGBM/SHAP cho giải thích đặc trưng, Perplexity cho tín hiệu xác suất, và Agent workflow để điều phối quá trình suy luận. Đây là cơ sở trực tiếp cho giải pháp đề xuất tại Chương 3.", 0),
    ]
    tables = [
        TableSpec(
            "2.1.1 Bài toán phát hiện mã nguồn do AI tạo ra",
            "Bảng 2.1: So sánh đặc trưng phong cách giữa mã do con người viết và mã do AI sinh ra",
            ["Đặc trưng", "Mã do con người viết", "Mã do AI sinh ra"],
            [
                ["Quy ước đặt tên biến", "Đa dạng, thường viết tắt hoặc theo thói quen cá nhân.", "Nhất quán, mô tả rõ ràng và gần với clean code."],
                ["Chú thích", "Ít, theo ngữ cảnh, đôi khi thiếu hoặc dùng ngôn ngữ bản địa.", "Có cấu trúc, giải thích mục đích hàm/lớp, thường bằng tiếng Anh chuẩn."],
                ["Định dạng code", "Không nhất quán, phụ thuộc IDE và thói quen.", "Chuẩn hóa cao, khoảng trắng và canh lề đồng nhất."],
                ["Cấu trúc logic", "Phản ánh tư duy cá nhân, đôi khi dồn logic vào một hàm.", "Có xu hướng module hóa và tách hàm rõ hơn."],
            ],
            [1.45, 2.45, 2.6],
        ),
        TableSpec(
            "2.1.3 Phân tích có tính giải thích",
            "Bảng 2.2: So sánh các phương pháp XAI liên quan đến phạm vi hiện tại",
            ["Phương pháp", "Loại", "Ưu điểm", "Hạn chế", "Trong hệ thống"],
            [
                ["SHAP", "Post-hoc, local/global", "Nhất quán, dựa trên Shapley values, phù hợp LightGBM.", "Chi phí tính toán cao hơn feature importance thường.", "Sử dụng"],
                ["LIME", "Post-hoc, local", "Dễ hiểu, model-agnostic.", "Kém ổn định giữa các lần chạy.", "Không sử dụng"],
                ["Attention weights", "Intrinsic", "Có sẵn trong Transformer.", "Còn tranh luận về khả năng giải thích thực sự.", "Không dùng làm XAI chính"],
                ["Perplexity", "Statistical signal", "Trực quan, hỗ trợ phát hiện mâu thuẫn.", "Không giải thích chi tiết từng đặc trưng.", "Tín hiệu bổ trợ"],
            ],
            [1.1, 1.25, 1.55, 1.5, 1.1],
        ),
        TableSpec(
            "2.2.4 Khoảng trống nghiên cứu",
            "Bảng 2.3: Tổng hợp so sánh các nghiên cứu liên quan",
            ["Nghiên cứu", "Năm", "Phương pháp", "Accuracy", "Ngôn ngữ", "XAI?"],
            [
                ["Kothari [2]", "2007", "Bayes + n-gram", "61%", "Đa ngôn ngữ", "Không"],
                ["Abuhamad [3]", "2019", "CNN + TF-IDF", "99.4%", "C++, Java, Python", "Không"],
                ["Abuhamad [4]", "2021", "LSTM + RF", "96.3%", "C++, Python", "Không"],
                ["Yang [10]", "2020", "n-gram + LogReg", "94.3%", "C++", "Không"],
                ["CLAVE [5][7]", "2025", "Contrastive + SVM", "92.3%", "Python", "Không"],
                ["ChatGPT Det. [15]", "2024", "DNN + RF + XGB", "98%", "Đa ngôn ngữ", "Feature-level"],
                ["Krish [1]", "2025", "Pre-trained + CNN", "-", "Đa ngôn ngữ", "Không"],
                ["Đề tài", "2026", "GraphCodeBERT + LightGBM/SHAP + Agent", "-", "C++", "Feature-level"],
            ],
            [1.25, 0.55, 1.65, 0.75, 1.05, 1.25],
        ),
    ]
    return "Chương 2 cập nhật v2", "Nền tảng lý thuyết và công trình liên quan", parts, tables


def ch3() -> tuple[str, str, list[tuple[str, int]], list[TableSpec]]:
    parts = [
        ("CHƯƠNG 3: GIẢI PHÁP ĐỀ XUẤT", 1),
        ("Chương này trình bày giải pháp đề xuất theo đúng phạm vi hiện tại của đề tài. So với bản gốc, các nội dung về thu thập dữ liệu, huấn luyện GraphCodeBERT, agent workflow, perplexity và giao diện phân tích vẫn được giữ vì là lõi của đồ án. Các phần thử nghiệm không còn là hướng chính, chẳng hạn ablation/mutation test tự động, được loại khỏi nội dung chính. Phần MLOps chỉ được trình bày ở mức hỗ trợ triển khai, không xem là đóng góp nghiên cứu trung tâm.", 0),
        ("3.1 Sơ đồ kiến trúc tổng thể", 2),
        ("3.1.1 Các giai đoạn chính của pipeline", 3),
        ("Hệ thống gồm bốn giai đoạn chính. Giai đoạn 1 là thu thập và tiền xử lý dữ liệu C++ từ GitHub, phân tách Basic/OOP và sinh phiên bản AI đối ngẫu bằng GPT-4o-mini. Giai đoạn 2 là huấn luyện mô hình phân loại GraphCodeBERT bằng sliding window và cross-validation. Giai đoạn 3 là suy luận bằng agent workflow gồm Router, Analyzer, Judge và Critique. Giai đoạn 4 là tích hợp ứng dụng web để người dùng nộp code, theo dõi tiến trình và xem kết quả phân tích.", 0),
        ("Ở phiên bản hiện tại, giao diện ứng dụng chính được định hướng theo Django thay vì chỉ Streamlit prototype. FastAPI giữ vai trò lớp điều phối AI, còn phần suy luận nặng có thể chạy trên GPU runtime riêng như Colab hoặc local GPU server. Cách tách này xuất phát từ ràng buộc tài nguyên: GraphCodeBERT ensemble và LLM phụ trợ cần GPU mạnh hơn máy phát triển thông thường.", 0),
        ("3.2 Quy trình xây dựng tập dữ liệu C++ mới", 2),
        ("3.2.1 Thu thập mã nguồn do con người viết", 3),
        ("Mã nguồn do con người viết được thu thập từ các repository công khai trên GitHub thông qua GitHub Search API. Nhóm xây dựng hai hướng thu thập tương ứng với Basic C++ và OOP C++. Tập Basic tập trung vào bài tập thuật toán, cấu trúc dữ liệu và lập trình cơ bản; tập OOP tập trung vào các dự án có class, kế thừa, đa hình và cấu trúc nhiều file.", 0),
        ("Các tiêu chí lọc chính gồm: repository được tạo và cập nhật trước thời điểm các LLM phổ biến rộng rãi, file có kích thước hợp lý, loại bỏ thư mục build/debug, loại bỏ repository dạng library/framework/engine/sdk và kiểm tra sự hiện diện của cấu trúc OOP đối với tập OOP. Điều kiện thời gian trước năm 2022 có vai trò quan trọng để giảm rủi ro dữ liệu human-written bị nhiễm mã do AI sinh ra.", 0),
        ("3.2.2 Tự động hóa sinh mã bằng AI", 3),
        ("Để tạo dữ liệu AI-generated, mỗi file human-written được đưa vào GPT-4o-mini với prompt yêu cầu tái cấu trúc mã nguồn theo phong cách AI nhưng giữ nguyên logic. Các thay đổi thường gồm đặt tên biến rõ nghĩa hơn, chuẩn hóa format, thêm chú thích và tách cấu trúc code gọn hơn. Cách tạo cặp dữ liệu đối ngẫu giúp giảm nhiễu từ nội dung thuật toán và làm nổi bật khác biệt phong cách.", 0),
        ("3.2.3 Tiền xử lý và nhãn hóa dữ liệu", 3),
        ("Dữ liệu sau khi thu thập và sinh mã được chuẩn hóa thành JSONL. Mỗi bản ghi chứa code, language, label, file_name, file_extension, relative_path, project_name, num_lines và num_chars. Hệ thống cũng dùng hash nội dung để phát hiện trùng lặp 100%, sau đó loại bỏ hoặc hợp nhất các mẫu trùng nhằm tránh rò rỉ dữ liệu giữa train/test.", 0),
        ("3.3 Phương pháp huấn luyện mô hình phân loại", 2),
        ("3.3.1 Kiến trúc mô hình", 3),
        ("Mô hình phân loại được xây dựng bằng cách fine-tune GraphCodeBERT (microsoft/graphcodebert-base) [23] với lớp RobertaForSequenceClassification cho hai nhãn AI và HUMAN. Việc chọn GraphCodeBERT được lý giải bởi khả năng khai thác thông tin data flow của mã nguồn, phù hợp với C++ và đặc biệt hữu ích khi phân biệt mã OOP với mã cơ bản.", 0),
        ("3.3.2 Sliding window cho mã nguồn dài", 3),
        ("Do Transformer có giới hạn đầu vào 512 token, các file dài được chia thành nhiều window. Mỗi window có tối đa 510 token nội dung, cộng với [CLS] và [SEP]. Các window chồng lấp 50%, tương ứng stride = 256, để giảm nguy cơ mất ngữ cảnh tại biên. Khi suy luận, xác suất AI của toàn file được tổng hợp từ các chunk bằng trung bình có trọng số theo độ dài chunk.", 0),
        ("3.3.3 Cấu hình huấn luyện", 3),
        ("Quá trình huấn luyện sử dụng learning rate 2e-5, batch size 32, 5 epochs, label smoothing 0.1, seed cố định và Stratified Group K-Fold. Group K-Fold đảm bảo các file thuộc cùng một project không bị tách sang nhiều fold khác nhau, qua đó giảm nguy cơ data leakage do các file cùng project thường chia sẻ phong cách lập trình.", 0),
        ("3.3.4 Kết hợp GraphCodeBERT và LightGBM", 3),
        ("Bên cạnh GraphCodeBERT, hệ thống sử dụng LightGBM trên các đặc trưng tĩnh của mã nguồn như comment_ratio, avg_identifier_length, brace_style_consistency, macro_count, maintainability_index, entropy và các chỉ số Halstead. LightGBM đóng vai trò nhánh phân tích hình thái, còn SHAP được dùng để giải thích đóng góp của từng đặc trưng. Việc kết hợp hai nhánh giúp hệ thống vừa có năng lực biểu diễn mã nguồn ở mức token, vừa có khả năng giải thích ở mức đặc trưng đọc được.", 0),
        ("3.4 Xây dựng hệ thống tác nhân thông minh", 2),
        ("3.4.1 Kiến trúc Agent Graph", 3),
        ("Agent workflow gồm bốn node chính. Router Node xác định code thuộc nhóm OOP hay NORMAL. Analyzer Node gọi mô hình phân loại và nhánh đặc trưng để lấy AI Score, LightGBM score và các tín hiệu hỗ trợ. Judge Node kiểm tra vùng nhập nhằng hoặc mâu thuẫn giữa AI Score và Perplexity. Critique Node tổng hợp kết quả thành nhận xét kỹ thuật cho người dùng.", 0),
        ("3.4.2 Router Node", 3),
        ("Router Node ưu tiên sử dụng LLM hoặc heuristic để phân loại OOP/NORMAL. Heuristic dựa trên các pattern như class, virtual, kế thừa, public/private/protected, template và toán tử phạm vi ClassName::method. Nếu tổng điểm vượt ngưỡng, code được định tuyến sang mô hình OOP; ngược lại dùng mô hình NORMAL.", 0),
        ("3.4.3 Analyzer Node", 3),
        ("Analyzer Node thực hiện ba nhiệm vụ. Thứ nhất, tokenize code và chia chunk theo sliding window. Thứ hai, gọi GraphCodeBERT runtime để lấy xác suất AI trên từng chunk và toàn file. Thứ ba, trích xuất đặc trưng tĩnh cho LightGBM và dùng SHAP để lấy các đặc trưng ảnh hưởng mạnh nhất.", 0),
        ("3.4.4 Judge Node", 3),
        ("Judge Node kiểm tra chất lượng kết quả theo các luật rõ ràng: điểm nằm trong vùng 0.40 đến 0.60 được xem là nhập nhằng; AI Score cao nhưng PPL cao có thể là mâu thuẫn; AI Score thấp nhưng PPL thấp cũng cần xem xét. Khi phát hiện mâu thuẫn và chưa retry, hệ thống có thể thử nhánh mô hình còn lại và chọn kết quả có độ tự tin cao hơn.", 0),
        ("3.4.5 Critique Node", 3),
        ("Critique Node dùng LLM để chuyển các tín hiệu kỹ thuật thành nhận xét dễ đọc. Prompt đầu vào gồm AI Score, model_used, Perplexity, top SHAP features, snippet đại diện và kết quả từng chunk. Đầu ra là phần global_critique và các nhận xét chi tiết giúp người dùng hiểu tại sao hệ thống nghiêng về AI hoặc HUMAN.", 0),
        ("3.5 Tích hợp giao diện ứng dụng", 2),
        ("3.5.1 Kiến trúc ứng dụng hiện tại", 3),
        ("Ở prototype ban đầu, giao diện được xây dựng bằng Streamlit và kết nối tới FastAPI trên Colab thông qua SSE. Ở hướng hiện tại của repo, Django được dùng làm web application chính để hỗ trợ đăng nhập, nộp code, lưu lịch sử và hiển thị dashboard. FastAPI đóng vai trò AI service nhận code_base64, chạy agent workflow và trả kết quả JSON hoặc SSE.", 0),
        ("3.5.2 Luồng xử lý yêu cầu", 3),
        ("Người dùng nộp code qua Django. Django lưu bản ghi đầu vào vào BronzeSubmission, mã hóa code thành Base64 và gọi FastAPI. FastAPI giải mã code, tính code_hash, kiểm tra cache, chạy AgentService, sau đó lưu kết quả phân tích vào Gold layer. Với chế độ streaming, Django proxy các SSE event từ FastAPI về trình duyệt để hiển thị tiến trình phân tích.", 0),
        ("3.5.3 Lưu trữ và dữ liệu vận hành ở mức vừa đủ", 3),
        ("Đồ án không đi sâu vào MLOps như một đóng góp độc lập. Tuy nhiên, hệ thống vẫn cần một cách tổ chức dữ liệu tối thiểu để lưu vết và tái sử dụng kết quả. Vì vậy, kiến trúc Bronze/Silver/Gold được dùng ở mức hỗ trợ: Bronze lưu code thô và metadata, Silver lưu đặc trưng/token phục vụ huấn luyện hoặc phân tích lại, Gold lưu kết quả dự đoán và giải thích phục vụ giao diện và dashboard.", 0),
        ("3.6 Tổng kết chương", 2),
        ("Chương 3 đã trình bày giải pháp đề xuất theo phạm vi đã được cập nhật. Lõi học thuật của đề tài gồm tập dữ liệu C++ tự xây dựng, mô hình GraphCodeBERT fine-tune, nhánh LightGBM/SHAP để giải thích đặc trưng, Perplexity làm tín hiệu hỗ trợ và Agent workflow để điều phối suy luận. Lớp ứng dụng gồm Django, FastAPI, GPU runtime và lưu trữ kết quả, được trình bày vừa đủ để chứng minh hệ thống có thể vận hành chứ không mở rộng thành một đề tài MLOps riêng.", 0),
    ]
    tables = [
        TableSpec(
            "3.2.3 Tiền xử lý và nhãn hóa dữ liệu",
            "Bảng 3.1: Tổng hợp các tập dữ liệu JSONL đã tạo",
            ["Tập dữ liệu", "File", "Kích thước", "Ghi chú"],
            [
                ["Basic C++ (ID)", "basic_cpp_ai_detection_dataset.jsonl", "~10.9 MB", "Dữ liệu mã C++ cơ bản (In-Distribution)."],
                ["OOP C++ (ID)", "OOP_cpp_ai_detection_dataset.jsonl", "~11.9 MB", "Dữ liệu mã hướng đối tượng (In-Distribution)."],
                ["Tập hợp nhất (ID)", "cpp_ai_detection_dataset.jsonl", "~22.8 MB", "Gộp Basic và OOP gốc phục vụ huấn luyện GraphCodeBERT."],
                ["Tập huấn luyện mở rộng (ML)", "train_data_with_oop.jsonl", "~18.8 MB", "Dữ liệu mở rộng gồm 6,880 mẫu ID và 2,000 mẫu OOD C/C++ để huấn luyện mô hình ML tĩnh."],
            ],
            [1.25, 2.2, 1.0, 2.05],
        ),
        TableSpec(
            "3.3.3 Cấu hình huấn luyện",
            "Bảng 3.2: Cấu hình siêu tham số huấn luyện GraphCodeBERT",
            ["Siêu tham số", "Giá trị", "Mô tả"],
            [
                ["model_name", "microsoft/graphcodebert-base", "Mô hình tiền huấn luyện."],
                ["max_len", "512", "Chiều dài tối đa đầu vào."],
                ["batch_size", "32", "Kích thước batch."],
                ["epochs", "5", "Số epoch huấn luyện."],
                ["learning rate", "2e-5", "Tốc độ học."],
                ["label_smoothing", "0.1", "Giảm overfitting."],
                ["n_folds", "5", "Số fold cross-validation."],
                ["stride", "256", "Sliding window overlap 50%."],
            ],
            [1.55, 1.55, 3.4],
        ),
        TableSpec(
            "3.4.2 Router Node",
            "Bảng 3.3: Hệ thống tính điểm OOP trong heuristic classifier",
            ["Từ khóa / Pattern", "Điểm"],
            [
                ["class + tên lớp", "+5"],
                ["virtual", "+3"],
                ["Kế thừa (: public/private/protected ClassName)", "+4"],
                ["public: / private: / protected:", "+2 mỗi loại"],
                ["template<", "+2"],
                ["ClassName::method", "+1, tối đa +3"],
            ],
            [4.8, 1.7],
        ),
        TableSpec(
            "3.5.3 Lưu trữ và dữ liệu vận hành ở mức vừa đủ",
            "Bảng 3.4: Vai trò dữ liệu trong hệ thống",
            ["Tầng", "Nội dung", "Mục đích"],
            [
                ["Bronze", "Code thô, metadata, code_hash.", "Lưu vết đầu vào và phục vụ kiểm tra lại."],
                ["Silver", "Đặc trưng ML và token DL.", "Phục vụ huấn luyện hoặc phân tích lại."],
                ["Gold", "Prediction, score, critique, SHAP output.", "Hiển thị kết quả và dashboard."],
            ],
            [1.0, 2.7, 2.8],
        ),
    ]
    return "Chương 3 cập nhật v2", "Giải pháp đề xuất", parts, tables


def build_doc(file_name: str, title: str, subtitle: str, parts, tables) -> Path:
    doc = Document()
    configure_doc(doc)
    add_title_page(doc, title, subtitle)

    remaining = list(tables)
    for text, level in parts:
        if level == 0:
            add_paragraph(doc, text)
        else:
            add_heading(doc, text, level)
            for table in [t for t in remaining if t.insert_after == text]:
                add_table(doc, table)
                remaining.remove(table)
    for table in remaining:
        add_table(doc, table)

    path = OUT_DIR / file_name
    doc.save(path)
    return path


def write_changelog(ch2_path: Path, ch3_path: Path) -> Path:
    path = OUT_DIR / "CHAPTER_CHANGELOG_v2.md"
    path.write_text(
        "\n".join(
            [
                "# Changelog cập nhật chương luận văn v2",
                "",
                "## Tệp được tạo",
                f"- `{ch2_path.name}`",
                f"- `{ch3_path.name}`",
                "",
                "## Nguyên tắc sửa",
                "- Dùng bản gốc làm xương sống học thuật thay vì bản rút gọn trước đó.",
                "- Giữ tác giả, năm, số tài liệu trong ngoặc vuông và công thức kỹ thuật.",
                "- Chỉ thêm các cập nhật thật sự cần để khớp với repo hiện tại.",
                "- Xóa LIG khỏi phần đóng góp chính; nếu cần nhắc, chỉ nhắc như nội dung đã loại khỏi phạm vi chính.",
                "- Viết phần MLOps ở mức hỗ trợ triển khai, không mở rộng thành trọng tâm luận văn.",
                "",
                "## Chương 2",
                "- Khôi phục mức độ kỹ thuật của Self-Attention với công thức `Attention(Q,K,V)`.",
                "- Khôi phục cách trích dẫn học thuật cho BERT, RoBERTa, CodeBERT, GraphCodeBERT, SHAP, Perplexity và CoT.",
                "- Chuyển phần XAI chính từ LIG sang SHAP/LightGBM theo phạm vi hiện tại.",
                "- Giữ literature review theo từng công trình cụ thể và bảng tổng hợp nghiên cứu.",
                "- Lập luận research gaps dựa trên bảng so sánh công trình, không chỉ phát biểu định tính.",
                "",
                "## Chương 3",
                "- Giữ các phần dataset, GraphCodeBERT training, sliding window, K-Fold, Router/Judge/Critique.",
                "- Bỏ mô tả mutation/ablation test như thành phần đang chạy chính.",
                "- Bỏ phần LIG khỏi pipeline chính; chuyển giải thích sang LightGBM/SHAP feature-level.",
                "- Cập nhật application layer từ Streamlit prototype sang Django + FastAPI + SSE theo repo.",
                "- Giữ Medallion/Celery/MLOps ở mức vừa đủ để mô tả lưu vết và vận hành.",
                "",
                "## Ghi chú QA",
                "- Đã kiểm tra cấu trúc DOCX bằng `python-docx`.",
                "- Chưa render PNG được vì môi trường không có `soffice`/LibreOffice.",
            ]
        ),
        encoding="utf-8",
    )
    return path


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ch2_title, ch2_subtitle, ch2_parts, ch2_tables = ch2()
    ch3_title, ch3_subtitle, ch3_parts, ch3_tables = ch3()
    ch2_path = build_doc("Chuong_2_cap_nhat_v2.docx", ch2_title, ch2_subtitle, ch2_parts, ch2_tables)
    ch3_path = build_doc("Chuong_3_cap_nhat_v2.docx", ch3_title, ch3_subtitle, ch3_parts, ch3_tables)
    write_changelog(ch2_path, ch3_path)


if __name__ == "__main__":
    main()
